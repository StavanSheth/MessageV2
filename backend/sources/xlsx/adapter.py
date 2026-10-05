import os
import re
from typing import List, Dict, Any, Tuple, Optional
import openpyxl
from backend.sources.base import SourceAdapter

class LocalXlsxSource(SourceAdapter):
    def __init__(self, file_path: str):
        self.file_path = file_path
        self.workbook = None
        self.sheet = None

    async def open(self) -> bool:
        if not os.path.exists(self.file_path):
            raise FileNotFoundError(f"File not found: {self.file_path}")
        self._last_mtime = os.path.getmtime(self.file_path)
        self.workbook = openpyxl.load_workbook(self.file_path, data_only=True)
        self.sheet = self.workbook.active
        return True

    async def validate_access(self) -> Tuple[bool, str]:
        if not os.path.exists(self.file_path):
            return False, "File does not exist on disk"
        try:
            wb = openpyxl.load_workbook(self.file_path, read_only=True)
            wb.close()
            return True, "File accessible"
        except Exception as e:
            return False, f"Access error: {str(e)}"

    @staticmethod
    def _normalize_header(header: Any) -> str:
        if not header:
            return ""
        return str(header).strip().lower().replace("_", " ").replace("-", " ")

    @staticmethod
    def _extract_username_from_url(url: str) -> Optional[str]:
        if not url:
            return None
        clean = str(url).split("?")[0].split("#")[0].rstrip("/")
        match = re.search(r"instagram\.com/([a-zA-Z0-9_\.\-]+)", clean, re.IGNORECASE)
        if match:
            user = match.group(1).strip()
            if user.lower() not in ["p", "reel", "reels", "stories", "explore", "direct"]:
                return user
        return None

    @staticmethod
    def _format_instagram_url(val: str) -> Tuple[str, str]:
        """Returns canonical (instagram_url, username)."""
        val_str = str(val).strip()
        clean = val_str.split("?")[0].split("#")[0].rstrip("/")
        if "instagram.com" in clean.lower():
            username = LocalXlsxSource._extract_username_from_url(clean) or ""
            if username:
                url = f"https://www.instagram.com/{username}/"
            else:
                url = clean if clean.startswith("http") else f"https://{clean}"
            return url, username
        else:
            clean_user = clean.lstrip("@").strip()
            if clean_user and re.match(r"^[a-zA-Z0-9_\.]+$", clean_user):
                username = clean_user
                url = f"https://www.instagram.com/{username}/"
            else:
                username = ""
                url = ""
            return url, username

    def _parse_sheet_rows(self, sheet_name: str, rows: list) -> Tuple[List[Dict[str, Any]], bool, Optional[str]]:
        if not rows:
            return [], False, "Empty sheet"

        # Find header row by scanning first 15 rows
        header_row_idx = None
        col_map = {}
        for r_i, r in enumerate(rows[:15]):
            test_map = {}
            for idx, col in enumerate(r):
                if col is None:
                    continue
                norm = self._normalize_header(col)
                # Ignore long paragraphs or sentences that are instructions, not column headers
                if norm and len(norm) <= 35 and len(norm.split()) <= 5:
                    test_map[norm] = idx

            # Must have at least 2 columns in a valid outreach table
            if len(test_map) < 2:
                continue

            has_ig = any(
                any(token in k.split() for token in ["instagram", "ig", "insta", "username", "handle"]) or
                k in ["instagram id/ link", "instagram id", "instagram link", "instagram url", "profile url", "ig handle"]
                for k in test_map
            )
            has_name = any(token in k.split() for token in ["name", "client", "contact", "lead", "user"] for k in test_map)

            if has_ig or (has_name and len(test_map) >= 3):
                header_row_idx = r_i
                col_map = test_map
                break

        if header_row_idx is None:
            return [], False, f"No recognized header row found in sheet '{sheet_name}'"

        # Determine column indexes with bounded alias matching
        def find_col(*aliases):
            for alias in aliases:
                for norm, idx in col_map.items():
                    if alias == norm or alias in norm.split():
                        return idx
                for norm, idx in col_map.items():
                    if len(norm) <= 30 and (alias in norm or norm in alias):
                        return idx
            return None

        name_idx = find_col("client name", "client", "full name", "lead name", "lead", "contact", "user", "name", "person")
        ig_idx = find_col(
            "instagram id/ link", "instagram id", "instagram link", "instagram url", "instagram profile",
            "ig handle", "ig username", "ig url", "profile url", "profile link", "instagram", "username",
            "ig link", "profile", "handle", "ig", "social link", "social", "account"
        )
        msg_idx = find_col("message", "dm", "text", "body", "initial message", "reason", "copy", "script", "first message")
        followers_idx = find_col("expected followers", "followers", "follower count", "following")
        notes_idx = find_col("notes", "note", "comment", "remarks", "details", "category", "industry")
        fu1_msg_idx = find_col("follow up 1 message", "follow-up 1 message", "followup 1 message", "fu1 message", "follow up 1", "follow-up 1", "fu1", "1st follow up", "touch 2")
        fu1_delay_idx = find_col("follow up 1 delay", "follow-up 1 delay", "followup 1 delay", "fu1 delay", "delay 1", "interval 1")
        fu2_msg_idx = find_col("follow up 2 message", "follow-up 2 message", "followup 2 message", "fu2 message", "follow up 2", "follow-up 2", "fu2", "2nd follow up", "touch 3")
        fu2_delay_idx = find_col("follow up 2 delay", "follow-up 2 delay", "followup 2 delay", "fu2 delay", "delay 2", "interval 2")
        replied_idx = find_col("replied status", "replied", "reply status", "has replied", "status")

        # If this sheet lacks an Instagram column, it is not an outreach lead sheet (e.g. Instructions, Summary)
        if ig_idx is None:
            return [], False, f"Sheet '{sheet_name}' lacks an Instagram column (checked {len(col_map)} headers)"

        records = []
        for row_idx, row in enumerate(rows[header_row_idx + 1:], start=header_row_idx + 2):
            # Skip empty rows
            if not any(cell is not None and str(cell).strip() != "" for cell in row):
                continue

            raw_dict = {}
            for col_norm, col_i in col_map.items():
                raw_dict[col_norm] = str(row[col_i]) if col_i < len(row) and row[col_i] is not None else ""

            name_val = str(row[name_idx]).strip() if name_idx is not None and name_idx < len(row) and row[name_idx] is not None else ""
            ig_val = str(row[ig_idx]).strip() if ig_idx is not None and ig_idx < len(row) and row[ig_idx] is not None else ""
            msg_val = str(row[msg_idx]).strip() if msg_idx is not None and msg_idx < len(row) and row[msg_idx] is not None else ""
            
            followers_val = None
            if followers_idx is not None and followers_idx < len(row) and row[followers_idx] is not None:
                try:
                    followers_val = int(re.sub(r"[^\d]", "", str(row[followers_idx])))
                except Exception:
                    followers_val = None

            notes_val = str(row[notes_idx]).strip() if notes_idx is not None and notes_idx < len(row) and row[notes_idx] is not None else None
            fu1_msg_val = str(row[fu1_msg_idx]).strip() if fu1_msg_idx is not None and fu1_msg_idx < len(row) and row[fu1_msg_idx] is not None else None
            fu2_msg_val = str(row[fu2_msg_idx]).strip() if fu2_msg_idx is not None and fu2_msg_idx < len(row) and row[fu2_msg_idx] is not None else None

            fu1_delay_val = 3
            if fu1_delay_idx is not None and fu1_delay_idx < len(row) and row[fu1_delay_idx] is not None:
                try:
                    fu1_delay_val = int(re.sub(r"[^\d]", "", str(row[fu1_delay_idx]))) or 3
                except Exception:
                    fu1_delay_val = 3

            fu2_delay_val = 5
            if fu2_delay_idx is not None and fu2_delay_idx < len(row) and row[fu2_delay_idx] is not None:
                try:
                    fu2_delay_val = int(re.sub(r"[^\d]", "", str(row[fu2_delay_idx]))) or 5
                except Exception:
                    fu2_delay_val = 5

            replied_val = "UNKNOWN"
            if replied_idx is not None and replied_idx < len(row) and row[replied_idx] is not None:
                r_str = str(row[replied_idx]).strip().upper()
                if r_str in ["YES", "TRUE", "Y"]:
                    replied_val = "YES"
                elif r_str in ["NO", "FALSE", "N"]:
                    replied_val = "NO"

            # Validation
            is_valid = True
            error_msg = None

            if not ig_val:
                is_valid = False
                error_msg = "Missing Instagram URL or username"
            else:
                ig_url, username = self._format_instagram_url(ig_val)
                if not username:
                    is_valid = False
                    error_msg = f"Invalid Instagram identifier: {ig_val}"

            if is_valid:
                if not name_val:
                    name_val = username or "User"
                if not msg_val:
                    msg_val = "Hey"

                norm_data = {
                    "name": name_val,
                    "instagram_url": ig_url,
                    "username": username,
                    "expected_followers": followers_val,
                    "message": msg_val,
                    "followup_1_message": fu1_msg_val,
                    "followup_1_delay_days": fu1_delay_val,
                    "followup_2_message": fu2_msg_val,
                    "followup_2_delay_days": fu2_delay_val,
                    "replied_status": replied_val,
                    "notes": notes_val
                }
            else:
                norm_data = {
                    "name": name_val or "Unknown",
                    "instagram_url": ig_val,
                    "username": None,
                    "message": msg_val or "Hey",
                    "followup_1_message": fu1_msg_val,
                    "followup_1_delay_days": fu1_delay_val,
                    "followup_2_message": fu2_msg_val,
                    "followup_2_delay_days": fu2_delay_val,
                    "replied_status": replied_val,
                    "notes": notes_val
                }

            records.append({
                "sheet_name": sheet_name,
                "row_number": row_idx,
                "raw": raw_dict,
                "normalized": norm_data,
                "is_valid": is_valid,
                "error": error_msg
            })

        return records, True, None

    async def read_records(self) -> List[Dict[str, Any]]:
        if not self.workbook:
            await self.open()

        all_records = []
        self.sheets_processed = []
        self.sheets_skipped = []

        worksheets = self.workbook.worksheets if hasattr(self.workbook, "worksheets") else [self.sheet]

        for ws in worksheets:
            sheet_rows = list(ws.iter_rows(values_only=True))
            if not sheet_rows:
                self.sheets_skipped.append({"sheet": ws.title, "reason": "Empty worksheet"})
                continue

            records, has_leads, reason = self._parse_sheet_rows(ws.title, sheet_rows)
            if has_leads:
                all_records.extend(records)
                self.sheets_processed.append({"sheet": ws.title, "records_count": len(records)})
            else:
                self.sheets_skipped.append({"sheet": ws.title, "reason": reason or "No lead columns found"})

        # If zero sheets contained leads, raise explicit Column Mismatch error
        if not all_records:
            sheet_names = [ws.title for ws in worksheets]
            raise ValueError(
                f"Column Mismatch Error: Could not find an Instagram identifier column (e.g. 'Instagram', 'Username', 'IG Handle') "
                f"in any worksheet. Checked sheets: {sheet_names}"
            )

        return all_records

    async def atomic_write_back(self, updates: Dict[str, Any], check_conflict: bool = True) -> Tuple[bool, str]:
        if not os.path.exists(self.file_path):
            return False, "File does not exist on disk"

        if check_conflict and hasattr(self, "_last_mtime") and self._last_mtime is not None:
            current_mtime = os.path.getmtime(self.file_path)
            if current_mtime > self._last_mtime + 0.5:
                return False, "SYNC_CONFLICT: Source file was modified externally"

        if self.workbook:
            try:
                self.workbook.close()
            except Exception:
                pass
            self.workbook = None

        try:
            wb = openpyxl.load_workbook(self.file_path)
            ws = wb.active
            headers = [cell.value for cell in ws[1]]
            status_col = None
            for idx, h in enumerate(headers, start=1):
                if h and str(h).strip().lower() in ["outreach status", "status"]:
                    status_col = idx
                    break
            if not status_col:
                status_col = len(headers) + 1
                ws.cell(row=1, column=status_col, value="Outreach Status")

            for key, data in updates.items():
                status_val = data.get("status", "")
                if key.startswith("row_"):
                    try:
                        row_num = int(key.split("_")[1])
                        ws.cell(row=row_num, column=status_col, value=status_val)
                    except Exception:
                        pass
                else:
                    for r_idx in range(2, ws.max_row + 1):
                        row_vals = [str(ws.cell(row=r_idx, column=c).value or "").lower() for c in range(1, 4)]
                        if any(key.lower() in rv for rv in row_vals):
                            ws.cell(row=r_idx, column=status_col, value=status_val)
                            break

            tmp_path = f"{self.file_path}.tmp"
            wb.save(tmp_path)
            wb.close()
            os.replace(tmp_path, self.file_path)
            self._last_mtime = os.path.getmtime(self.file_path)
            return True, "Write-back successful"
        except Exception as e:
            return False, str(e)

    async def update_record(self, record_id: str, data: Dict[str, Any]) -> bool:
        if self.workbook:
            try:
                self.workbook.close()
            except Exception:
                pass
            self.workbook = None

        try:
            wb = openpyxl.load_workbook(self.file_path)
            ws = wb.active
            headers = [cell.value for cell in ws[1]]
            status_col = None
            for idx, h in enumerate(headers, start=1):
                if h and str(h).strip().lower() in ["outreach status", "status"]:
                    status_col = idx
                    break
            if not status_col:
                status_col = len(headers) + 1
                ws.cell(row=1, column=status_col, value="Outreach Status")

            status_val = data.get("status", "")
            target_row = None

            if str(record_id).startswith("row_"):
                try:
                    target_row = int(str(record_id).split("_")[1])
                except Exception:
                    pass
            else:
                clean_rec = str(record_id).lower().lstrip("@")
                for r_idx in range(2, ws.max_row + 1):
                    row_vals = [str(ws.cell(row=r_idx, column=c).value or "").lower() for c in range(1, 5)]
                    if any(clean_rec == rv or clean_rec in rv for rv in row_vals):
                        target_row = r_idx
                        break

            if target_row and target_row <= ws.max_row:
                ws.cell(row=target_row, column=status_col, value=status_val)
                wb.save(self.file_path)
                wb.close()
                return True

            wb.close()
            return False
        except Exception:
            return False

    async def sync(self) -> Dict[str, Any]:
        records = await self.read_records()
        valid = sum(1 for r in records if r["is_valid"])
        return {
            "total": len(records),
            "valid": valid,
            "invalid": len(records) - valid,
            "records": records,
            "sheets_processed": getattr(self, "sheets_processed", []),
            "sheets_skipped": getattr(self, "sheets_skipped", [])
        }

    async def close(self) -> None:
        if self.workbook:
            self.workbook.close()
            self.workbook = None
