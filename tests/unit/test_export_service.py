import pytest
import openpyxl
import io
from httpx import AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession
from datetime import datetime, timezone

from backend.database.models import Contact, Task
from backend.services.export_service import ExportService

@pytest.mark.asyncio
async def test_export_service_workbooks(test_session: AsyncSession):
    # Setup sample data in isolated test session
    c1 = Contact(
        name="Lead One",
        username="lead_one",
        instagram_url="https://instagram.com/lead_one",
        message="Hey! Custom outreach 1",
        replied_status="NO"
    )
    c2 = Contact(
        name="Lead Replied",
        username="lead_replied",
        instagram_url="https://instagram.com/lead_replied",
        replied_status="YES",
        extracted_phone="+1234567890"
    )
    test_session.add_all([c1, c2])
    await test_session.commit()
    await test_session.refresh(c1)
    await test_session.refresh(c2)

    t1 = Task(contact_id=c1.id, type="MESSAGE", status="READY", scheduled_at=datetime.now(timezone.utc))
    t2 = Task(contact_id=c2.id, type="FOLLOW_UP_1", status="COMPLETED", completed_at=datetime.now(timezone.utc))
    t3 = Task(contact_id=c2.id, type="FOLLOW_UP_2", status="MANUAL_REVIEW", manual_review_reason="[EXTERNAL_MESSAGE_DETECTED] External message detected")
    test_session.add_all([t1, t2, t3])
    await test_session.commit()

    # 1. Test generate_universal_excel
    u_stream = await ExportService.generate_universal_excel(test_session)
    u_wb = openpyxl.load_workbook(u_stream)
    assert len(u_wb.sheetnames) == 5
    assert "1. Contacts Directory" in u_wb.sheetnames
    assert "2. Dispatch Queue" in u_wb.sheetnames

    # 2. Test generate_queue_excel
    q_stream = await ExportService.generate_queue_excel(test_session)
    q_wb = openpyxl.load_workbook(q_stream)
    assert len(q_wb.sheetnames) == 4
    assert q_wb.sheetnames == ["Upcoming Queue", "Done Already", "Action Needed", "All Tasks"]
    # Check that Action Needed sheet has t3
    ws_action = q_wb["Action Needed"]
    assert ws_action.max_row >= 2

    # 3. Test generate_contacts_excel
    c_stream = await ExportService.generate_contacts_excel(test_session)
    c_wb = openpyxl.load_workbook(c_stream)
    assert len(c_wb.sheetnames) == 4
    assert c_wb.sheetnames == ["All Contacts", "Replied Leads", "Active Outreach", "Needs Review"]
    # Check that Replied Leads sheet has c2
    ws_replied = c_wb["Replied Leads"]
    assert ws_replied.max_row >= 2


@pytest.mark.asyncio
async def test_export_endpoints_cache_headers(client: AsyncClient):
    # Test tasks export
    res_q = await client.get("/api/tasks/export/excel")
    assert res_q.status_code == 200
    assert "no-cache" in res_q.headers.get("Cache-Control", "")
    assert res_q.headers.get("Content-Disposition", "").startswith("attachment; filename=Dispatch_Queue_")

    # Test contacts export
    res_c = await client.get("/api/contacts/export/excel")
    assert res_c.status_code == 200
    assert "no-cache" in res_c.headers.get("Cache-Control", "")
    assert res_c.headers.get("Content-Disposition", "").startswith("attachment; filename=Instagram_Outreach_Tracking_")

    # Test universal export
    res_u = await client.get("/api/sources/export/universal")
    assert res_u.status_code == 200
    assert "no-cache" in res_u.headers.get("Cache-Control", "")
    assert res_u.headers.get("Content-Disposition", "").startswith("attachment; filename=Universal_Instagram_Outreach_")
