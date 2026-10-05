import pytest
import openpyxl
from backend.sources.xlsx.adapter import LocalXlsxSource

def test_extract_username_variations():
    assert LocalXlsxSource._extract_username_from_url("https://www.instagram.com/john_doe/") == "john_doe"
    assert LocalXlsxSource._extract_username_from_url("https://instagram.com/jane.doe") == "jane.doe"
    
    url, user = LocalXlsxSource._format_instagram_url("@my_handle")
    assert user == "my_handle"
    assert url == "https://www.instagram.com/my_handle/"

    url, user = LocalXlsxSource._format_instagram_url("simple_user")
    assert user == "simple_user"
    assert url == "https://www.instagram.com/simple_user/"

def test_header_normalization():
    assert LocalXlsxSource._normalize_header("  IG_Profile  ") == "ig profile"
    assert LocalXlsxSource._normalize_header("Full-Name") == "full name"

@pytest.mark.asyncio
async def test_parse_in_memory_xlsx(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contacts"
    ws.append(["Name", "Instagram URL", "Message"])
    ws.append(["Alice Smith", "https://instagram.com/alice_smith", "Hello Alice!"])
    ws.append(["Bob Jones", "@bobjones", "Hey Bob"])

    file_path = tmp_path / "test_contacts.xlsx"
    wb.save(file_path)

    adapter = LocalXlsxSource(file_path=str(file_path))
    records = await adapter.read_records()
    assert len(records) == 2
    assert records[0]["normalized"]["username"] == "alice_smith"
    assert records[0]["normalized"]["name"] == "Alice Smith"
    assert records[0]["normalized"]["message"] == "Hello Alice!"
    assert records[1]["normalized"]["username"] == "bobjones"

@pytest.mark.asyncio
async def test_multisheet_filters_irrelevant_data(tmp_path):
    wb = openpyxl.Workbook()
    # Sheet 1: Irrelevant Instructions / Notes
    ws_notes = wb.active
    ws_notes.title = "Instructions"
    ws_notes.append(["How to Use This Template"])
    ws_notes.append(["Please enter lead handles in the Leads tab below."])
    ws_notes.append(["Do not edit columns."])

    # Sheet 2: Real Leads
    ws_leads = wb.create_sheet(title="October Leads")
    ws_leads.append(["Lead Name", "IG Handle", "Notes"])
    ws_leads.append(["Dave Miller", "@davemiller", "Priority"])
    ws_leads.append(["Eve Adams", "https://instagram.com/eveadams", "VIP"])

    # Sheet 3: Summary / Empty
    ws_summary = wb.create_sheet(title="Dashboard Summary")
    ws_summary.append(["Metric", "Count"])
    ws_summary.append(["Total Sent", 0])

    file_path = tmp_path / "multisheet_leads.xlsx"
    wb.save(file_path)

    adapter = LocalXlsxSource(file_path=str(file_path))
    result = await adapter.sync()
    
    assert result["total"] == 2
    assert result["valid"] == 2
    assert len(result["sheets_processed"]) == 1
    assert result["sheets_processed"][0]["sheet"] == "October Leads"
    assert len(result["sheets_skipped"]) == 2

@pytest.mark.asyncio
async def test_column_mismatch_raises_informative_error(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RandomData"
    ws.append(["First Name", "City", "Company", "Phone"])
    ws.append(["John", "New York", "Acme", "555-1234"])

    file_path = tmp_path / "missing_ig_column.xlsx"
    wb.save(file_path)

    adapter = LocalXlsxSource(file_path=str(file_path))
    with pytest.raises(ValueError) as exc_info:
        await adapter.read_records()
    assert "Column Mismatch Error" in str(exc_info.value)
    assert "RandomData" in str(exc_info.value)

